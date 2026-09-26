"""Generate fake bank statements for testing Money Leak Finder.

Every name, amount and account number here is invented.
The answer key for these files is tests/answer-key.json (written by hand).

    pip install openpyxl
    python3 tests/make-samples.py
"""
import csv
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

import os

OUT = Path(os.environ.get("SAMPLES_OUT") or Path(__file__).resolve().parent.parent / "samples")
OUT.mkdir(parents=True, exist_ok=True)
# SEED changes the random everyday spending (used by tests/stress.sh); the
# subscriptions, duplicates and one-offs in the answer key stay the same.
SEED = int(os.environ.get("SEED", "0"))


def monthly(start, end, day):
    """Dates on `day` of every month between start and end (inclusive)."""
    out = []
    y, m = start.year, start.month
    while True:
        d = date(y, m, day)
        if d > end:
            break
        if d >= start:
            out.append(d)
        m += 1
        if m == 13:
            y, m = y + 1, 1
    return out


def next_weekday(d):
    """UK direct debits move to the next working day when they fall on a weekend."""
    while d.weekday() >= 5:
        d += timedelta(days=1)
    return d


def every(start, end, step_days):
    out, d = [], start
    while d <= end:
        out.append(d)
        d += timedelta(days=step_days)
    return out


def random_dates(rng, start, end, min_gap, max_gap):
    out, d = [], start + timedelta(days=rng.randint(0, max_gap))
    while d <= end:
        out.append(d)
        d += timedelta(days=rng.randint(min_gap, max_gap))
    return out


def unique_amount(rng, lo, hi, used):
    while True:
        a = round(rng.uniform(lo, hi), 2)
        if a not in used:
            used.add(a)
            return a


# ---------------------------------------------------------------------------
# Sample 1: UK current account, CSV
#   DD/MM/YYYY, separate "Paid out" / "Paid in" columns, £ signs, commas in
#   numbers, a balance column, 4 junk rows above the header, newest first.
# ---------------------------------------------------------------------------
def sample_uk():
    rng = random.Random(1 + SEED * 10)
    start, end = date(2025, 1, 1), date(2025, 6, 30)
    tx = []  # (date, description, out, in)

    def out(d, desc, amt):
        tx.append((d, desc, amt, None))

    def inc(d, desc, amt):
        tx.append((d, desc, None, amt))

    for d in monthly(start, end, 1):
        out(d, "STO J PATEL RENT", 950.00)
        out(next_weekday(d), "DD PUREGYM LTD", 24.99)
    for d in monthly(start, end, 15):
        out(next_weekday(d), "DD CROYDON COUNCIL 88123456", 142.00)
    for d in monthly(start, end, 3):  # price rise in April
        out(d, f"CARD PAYMENT TO NETFLIX.COM ON {d:%d/%m}", 10.99 if d.month < 4 else 12.99)
    for d in monthly(start, end, 12):
        ref = "".join(rng.choice("ABCDEFGHJKLMNPQRSTUVWXYZ0123456789") for _ in range(10))
        if sum(c.isdigit() for c in ref) < 3:
            ref = ref[:7] + "482"
        out(next_weekday(d), f"DD SPOTIFY {ref}", 11.99)
    for d in monthly(start, end, 20):
        out(next_weekday(d), "DD EE LIMITED", 32.00)
    for d in monthly(start, end, 8):
        out(d, f"CARD PAYMENT TO AMAZON PRIME*{rng.randint(10**8, 10**9)}X ON {d:%d/%m}", 8.99)
    for d in monthly(start, date(2025, 3, 31), 10):  # cancelled after March
        out(next_weekday(d), "DD DISNEY PLUS", 7.99)
    for d in monthly(start, end, 25):
        inc(next_weekday(d), "BGC ACME LTD SALARY", 2450.00)

    # Duplicate: the same takeaway order charged twice, one day apart
    out(date(2025, 3, 14), "CARD PAYMENT TO DELIVEROO ON 14/03", 23.45)
    out(date(2025, 3, 15), "CARD PAYMENT TO DELIVEROO ON 15/03", 23.45)
    for d, a in [(date(2025, 2, 2), 18.20), (date(2025, 4, 21), 27.10), (date(2025, 5, 30), 15.85)]:
        out(d, f"CARD PAYMENT TO DELIVEROO ON {d:%d/%m}", a)

    # One-off purchases: must NOT be flagged
    out(date(2025, 2, 22), "CARD PAYMENT TO CURRYS PC WORLD ON 22/02", 349.00)
    out(date(2025, 5, 9), "CARD PAYMENT TO JOHN LEWIS ON 09/05", 89.00)
    out(date(2025, 4, 17), "CARD PAYMENT TO TRAINLINE ON 17/04", 64.30)

    # Everyday spending: must NOT be flagged
    used = set()
    for d in every(date(2025, 1, 4), end, 7):
        out(d, f"CARD PAYMENT TO TESCO STORES 3297 ON {d:%d/%m}", unique_amount(rng, 25, 95, used))
    for d in every(date(2025, 1, 7), end, 7):  # Tuesdays and Thursdays: same coffee
        out(d, f"CARD PAYMENT TO PRET A MANGER ON {d:%d/%m}", 3.95)
        d2 = d + timedelta(days=2)
        if d2 <= end:
            out(d2, f"CARD PAYMENT TO PRET A MANGER ON {d2:%d/%m}", 3.95)
    for d in random_dates(rng, start, end, 11, 24):
        out(d, f"CARD PAYMENT TO SHELL CROYDON ON {d:%d/%m}", unique_amount(rng, 45, 72, used))

    # Oldest first to compute the running balance, then write newest first
    tx.sort(key=lambda t: (t[0], t[1]))
    bal = 3200.00
    rows = []
    for d, desc, o, i in tx:
        bal += (i or 0) - (o or 0)
        rows.append((d, desc, o, i, bal))
    rows.reverse()

    money = lambda v: "" if v is None else f"£{v:,.2f}"
    with open(OUT / "uk-bank-statement.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
        w.writerow(["Account Name:", "Current Account - J SMITH"])
        w.writerow(["Account Number:", "20-45-67 XXXX4821"])
        w.writerow(["Statement Period:", "01/01/2025 to 30/06/2025"])
        w.writerow([])
        w.writerow(["Date", "Description", "Paid out", "Paid in", "Balance"])
        for d, desc, o, i, b in rows:
            w.writerow([f"{d:%d/%m/%Y}", desc, money(o), money(i), money(b)])


# ---------------------------------------------------------------------------
# Sample 2: US credit card, CSV
#   MM/DD/YYYY, one Amount column where charges are POSITIVE and payments
#   are negative, "$" signs, commas in numbers, locations after long spaces.
# ---------------------------------------------------------------------------
def sample_us():
    rng = random.Random(2 + SEED * 10)
    start, end = date(2024, 7, 1), date(2025, 7, 31)
    tx = []

    def ch(d, desc, amt):
        tx.append((d, desc, amt))

    for n, d in enumerate(monthly(start, end, 5)):
        desc = "NETFLIX.COM          LOS GATOS           CA" if n % 2 else "NETFLIX.COM 866-579-7172 CA"
        ch(d, desc, 15.49 if d < date(2025, 1, 1) else 17.99)
    for d in monthly(start, end, 18):
        ch(d, "SPOTIFY USA          877-778-1161        NY", 11.99)
    for d in monthly(start, end, 22):
        ch(d, "APPLE.COM/BILL       866-712-7753        CA", 2.99)
    ch(date(2025, 3, 9), "APPLE.COM/BILL       866-712-7753        CA", 49.99)  # one-off app purchase
    for d in every(date(2025, 4, 7), end, 7):
        if d != date(2025, 6, 2):  # skipped one week
            ch(d, "HELLOFRESH           NEW YORK            NY", 69.99)
    for d in [date(2024, 7, 10), date(2024, 10, 10), date(2025, 1, 10), date(2025, 4, 10), date(2025, 7, 10)]:
        ch(d, "STATE FARM INSURANCE 800-956-6310        IL", 186.00)
    ch(date(2024, 7, 14), "AMAZON PRIME*RT4K29XL0 AMZN.COM/BILL   WA", 139.00)
    ch(date(2025, 7, 14), "AMAZON PRIME*8H2QP71M3 AMZN.COM/BILL   WA", 139.00)
    for d in monthly(start, end, 2):
        ch(d, "COMCAST CABLE COMM   800-266-2278        PA", 89.99)
    ch(date(2025, 3, 3), "COMCAST CABLE COMM   800-266-2278        PA", 89.99)  # duplicate
    for d in monthly(start, end, 17):
        ch(d, "PLANET FITNESS       HAMPTON             NH", 10.00)
    ch(date(2025, 1, 15), "PLANET FITNESS       HAMPTON             NH", 49.00)  # annual fee, one-off

    # Charged twice but refunded: must NOT be flagged
    ch(date(2025, 5, 20), "AMAZON MKTPLACE PMTS AMZN.COM/BILL       WA", 34.99)
    ch(date(2025, 5, 20), "AMAZON MKTPLACE PMTS AMZN.COM/BILL       WA", 34.99)
    ch(date(2025, 5, 27), "AMAZON MKTPLACE PMTS AMZN.COM/BILL       WA", -34.99)

    # One-offs
    ch(date(2024, 11, 29), "BEST BUY 00012345     SAN JOSE            CA", 1299.00)
    ch(date(2025, 3, 21), "DELTA AIR 0062345678901 ATLANTA          GA", 412.60)

    used = {34.99}
    for d in random_dates(rng, start, end, 5, 30):
        ch(d, "AMAZON MKTPLACE PMTS AMZN.COM/BILL       WA", unique_amount(rng, 9, 120, used))
    for d in random_dates(rng, start, end, 2, 9):
        ch(d, "STARBUCKS STORE 12345 SEATTLE            WA", unique_amount(rng, 4.25, 8.95, used))
    for d in random_dates(rng, start, end, 5, 9):
        ch(d, "WHOLEFDS MKT 10234   AUSTIN              TX", unique_amount(rng, 40, 160, used))
    for d in random_dates(rng, start, end, 9, 18):
        ch(d, "SHELL OIL 57442153   AUSTIN              TX", unique_amount(rng, 35, 60, used))

    for d in monthly(start, end, 25):
        ch(d, "AUTOPAY PAYMENT - THANK YOU", -round(rng.uniform(1200, 2400), 2))

    tx.sort(key=lambda t: (t[0], t[1]), reverse=True)

    def money(v):
        return f"-${-v:,.2f}" if v < 0 else f"${v:,.2f}"

    with open(OUT / "us-credit-card.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, quoting=csv.QUOTE_ALL)
        w.writerow(["Date", "Description", "Card Member", "Account #", "Amount"])
        for d, desc, a in tx:
            w.writerow([f"{d:%m/%d/%Y}", desc, "JANE DOE", "-41007", money(a)])


# ---------------------------------------------------------------------------
# Sample 3: European bank export, Excel .xlsx
#   Real Excel date cells (dd-mm-yyyy), amounts always positive with a
#   separate "Debit/credit" column, title rows above the header, newest first.
# ---------------------------------------------------------------------------
def sample_eu():
    rng = random.Random(3 + SEED * 10)
    start, end = date(2024, 10, 1), date(2025, 3, 31)
    tx = []  # (date, name, dc, amount, code, type, notes)

    def debit(d, name, amt, code="IC", kind="Direct debit", notes=""):
        tx.append((d, name, "Debit", amt, code, kind, notes))

    def credit(d, name, amt, notes=""):
        tx.append((d, name, "Credit", amt, "OV", "Transfer", notes))

    def dd_note(name, d):
        return (f"Name: {name} Description: Klantnummer {rng.randint(10**7, 10**8)} "
                f"IBAN: NL{rng.randint(10, 99)}ABNA0{rng.randint(10**8, 10**9)} "
                f"Mandate ID: {rng.randint(10**6, 10**7)} Date/time: {d:%d-%m-%Y}")

    def card_note(d):
        return f"Card sequence no.: 007 {d:%d-%m-%Y} {rng.randint(8, 21)}:{rng.randint(10, 59)} Transaction: {rng.randint(10**5, 10**6):X} Term: {rng.randint(10**7, 10**8)}"

    for d in monthly(start, end, 27):
        debit(d, "Ziggo Services BV", 54.00, notes=dd_note("Ziggo Services BV", d))
    for d, a in zip(monthly(start, end, 21), [98.00, 105.00, 112.00, 121.00, 118.00, 101.00]):
        debit(d, "Vattenfall Klantenservice N.V.", a, notes=dd_note("Vattenfall", d))
    for d in monthly(start, end, 7):
        debit(d, "Spotify AB", 10.99 if d.year == 2024 else 11.99, notes=dd_note("Spotify AB", d))
    for d in every(date(2024, 10, 3), end, 28):
        debit(d, "Basic-Fit Nederland BV", 29.99, notes=dd_note("Basic-Fit", d))
    for d in [date(2024, 10, 1), date(2025, 1, 2)]:
        debit(d, "Centraal Beheer", 87.50, notes=dd_note("Centraal Beheer", d))
    for d in [date(2024, 10, 15), date(2024, 11, 15), date(2024, 12, 15)]:
        debit(d, "Videoland", 9.99, notes=dd_note("Videoland", d))
    for d in monthly(start, end, 24):
        credit(d, "Acme Nederland BV", 3150.00, notes="Salaris")

    # Duplicate order
    debit(date(2024, 11, 12), "bol.com", 59.95, code="ID", kind="iDEAL", notes="Bestelling 4012345678")
    debit(date(2024, 11, 13), "bol.com", 59.95, code="ID", kind="iDEAL", notes="Bestelling 4012345679")
    debit(date(2024, 12, 5), "bol.com", 24.99, code="ID", kind="iDEAL", notes="Bestelling 4012399911")
    debit(date(2025, 2, 18), "bol.com", 12.50, code="ID", kind="iDEAL", notes="Bestelling 4012456002")

    # One-offs
    debit(date(2024, 11, 29), "MediaMarkt Amsterdam", 899.00, code="BA", kind="Payment terminal", notes=card_note(date(2024, 11, 29)))
    debit(date(2025, 1, 20), "KLM Royal Dutch Airlines", 245.00, code="ID", kind="iDEAL", notes="Booking ABC123")

    used = set()
    for d in random_dates(rng, start, end, 2, 5):
        debit(d, "Albert Heijn 1234", unique_amount(rng, 8, 75, used), code="BA", kind="Payment terminal", notes=card_note(d))
    for d in random_dates(rng, start, end, 5, 10):
        debit(d, "Jumbo Amsterdam", unique_amount(rng, 10, 60, used), code="BA", kind="Payment terminal", notes=card_note(d))
    for d in every(date(2024, 10, 1), end, 7):  # commuting: same fare twice a week
        for dd in (d, d + timedelta(days=3)):
            if dd <= end:
                debit(dd, "NS GROEP IZ NS REIZIGERS", 4.20, code="BA", kind="Payment terminal", notes=card_note(dd))

    tx.sort(key=lambda t: (t[0], t[1]), reverse=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(["ING - Transactions export"])
    ws.append(["Account: NL91 INGB 0001 2345 67 - Period 01-10-2024 to 31-03-2025"])
    ws.append([])
    ws.append(["Date", "Name / Description", "Account", "Counterparty", "Code", "Debit/credit",
               "Amount (EUR)", "Transaction type", "Notifications"])
    for d, name, dc, amt, code, kind, notes in tx:
        ws.append([d, name, "NL91INGB0001234567",
                   "" if code == "BA" else f"NL{rng.randint(10, 99)}RABO0{rng.randint(10**8, 10**9)}",
                   code, dc, amt, kind, notes])
        r = ws.max_row
        ws.cell(row=r, column=1).number_format = "dd-mm-yyyy"
        ws.cell(row=r, column=7).number_format = "#,##0.00"
    # A second, unrelated sheet the reader must ignore
    ws2 = wb.create_sheet("Notes")
    ws2.append(["Exported from online banking"])
    wb.save(OUT / "eu-bank-export.xlsx")


# ---------------------------------------------------------------------------
# Extra edge case: no header row, semicolons, decimal commas (1.299,00),
# DD.MM.YYYY dates, oldest first. The tool should ask which column is which.
# ---------------------------------------------------------------------------
def sample_no_header():
    rng = random.Random(4 + SEED * 10)
    start, end = date(2025, 1, 1), date(2025, 4, 30)
    tx = []
    for d in monthly(start, end, 3):
        tx.append((d, f"LASTSCHRIFT NETFLIX INTERNATIONAL B.V. REF {rng.randint(10**9, 10**10)}", -13.99))
    for d in monthly(start, end, 5):
        tx.append((d, f"LASTSCHRIFT TELEKOM DEUTSCHLAND GMBH RECHNUNG {rng.randint(10**9, 10**10)}", -39.95))
    for d in monthly(start, end, 28):
        tx.append((d, "GEHALT ACME GMBH", 3200.00))
    tx.append((date(2025, 2, 14), "KARTENZAHLUNG SATURN ELECTRO BERLIN 14.02.2025", -1299.00))
    used = set()
    for d in random_dates(rng, start, end, 3, 7):
        tx.append((d, f"KARTENZAHLUNG REWE MARKT GMBH {d:%d.%m.%Y}", -unique_amount(rng, 12, 90, used)))
    tx.sort(key=lambda t: (t[0], t[1]))

    def money(v):
        s = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        return ("-" if v < 0 else "") + s

    with open(OUT / "no-header-semicolon.csv", "w", newline="", encoding="utf-8") as f:
        for d, desc, a in tx:
            f.write(f"{d:%d.%m.%Y};{desc};{money(a)};EUR\n")


if __name__ == "__main__":
    sample_uk()
    sample_us()
    sample_eu()
    sample_no_header()
    print("Wrote samples to", OUT)
